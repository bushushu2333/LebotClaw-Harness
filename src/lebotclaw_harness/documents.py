"""Editable office artifacts. Content and structure come from the agent."""
import io
from .artifacts import save_artifact


def document_create(args, ctx):
    path = ctx.path(args['path'])
    if path.exists():
        raise ValueError('请使用新文件名保存文档版本。')
    data = io.BytesIO()
    if path.suffix.lower() == '.docx':
        from docx import Document
        from docx.shared import Inches
        doc = Document()
        doc.add_heading(args['title'], 0)
        for section in args['sections']:
            doc.add_heading(section['title'], 1)
            for paragraph in section.get('body', '').split('\n'):
                if paragraph: doc.add_paragraph(paragraph)
            for item in section.get('bullets', []): doc.add_paragraph(item, style='List Bullet')
            if section.get('image'):
                doc.add_picture(str(ctx.path(section['image'])), width=Inches(5.8))
        doc.save(data)
    elif path.suffix.lower() == '.pptx':
        from pptx import Presentation
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
        sections = [{'title':args['title'], 'body':args.get('subtitle', '')}, *args['sections']]
        for index, section in enumerate(sections):
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            slide.background.fill.solid(); slide.background.fill.fore_color.rgb = RGBColor.from_string('F5F4FA')
            def box(x, y, w, h, text, size, color='222239', bold=False):
                frame = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame
                frame.word_wrap = True
                frame.text = text
                for paragraph in frame.paragraphs:
                    paragraph.font.name = 'Microsoft YaHei'
                    paragraph.font.size = Pt(size)
                    paragraph.font.bold = bold
                    paragraph.font.color.rgb = RGBColor.from_string(color)
                return frame
            box(.7, .65, 11.9, 1.2, section['title'], 32, '6657F4', True)
            text = section.get('body', '')
            if section.get('bullets'): text += '\n' + '\n'.join('• ' + s for s in section['bullets'])
            box(.75, 2, 6 if section.get('image') else 11.7, 4.4, text.strip(), 22)
            if section.get('image'):
                from PIL import Image
                image = ctx.path(section['image'])
                with Image.open(image) as im: ratio = im.width / im.height
                w = min(4.6, 3.9 * ratio); h = w / ratio
                slide.shapes.add_picture(str(image), Inches(7.85), Inches(2.05), width=Inches(w), height=Inches(h))
            box(.75, 6.95, 11, .3, '超级小博 · ' + args['title'], 11, '62647A')
            box(12, 6.9, .7, .4, str(index + 1), 12, '62647A')
        prs.save(data)
    else:
        raise ValueError('文档输出支持 .docx 或 .pptx。')
    result = save_artifact(ctx, args['path'], data.getvalue())
    return {**result, 'editable':True, 'sections':len(args['sections']), 'visual_check':'not_performed'}


def office_read(path):
    if path.stat().st_size > 15_000_000:
        raise ValueError('文档最多 15 MB。')
    # Zip-bomb guard before office parsers expand XML.
    import zipfile
    with zipfile.ZipFile(path) as z:
        if sum(i.file_size for i in z.infolist()) > 80_000_000:
            raise ValueError('文档解压后过大。')
    if path.suffix.lower() == '.docx':
        from docx import Document
        doc = Document(str(path))
        text = '\n'.join(p.text for p in doc.paragraphs)
        text += '\n' + '\n'.join(' | '.join(c.text for c in row.cells) for t in doc.tables for row in t.rows)
    else:
        from pptx import Presentation
        doc = Presentation(str(path))
        text = '\n'.join('第 %s 页\n%s' % (i+1, '\n'.join(s.text for s in slide.shapes if s.has_text_frame)) for i, slide in enumerate(doc.slides))
    return {'text':text[:40000], 'truncated':len(text)>40000}
