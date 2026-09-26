import asyncio
import base64
import io
import json
from pathlib import Path
import httpx
import pytest
from PIL import Image
from lebotclaw_harness.config import Config
from lebotclaw_harness.runtime import Runtime
from lebotclaw_harness.tools import ToolContext, default_registry
from lebotclaw_harness.preview import PreviewHost, browser_check
from lebotclaw_harness.model import ModelReply
from conftest import tool_reply, sse
from test_runtime import configured


def png():
    b=io.BytesIO(); Image.new('RGB',(48,32),(100,80,245)).save(b,format='PNG'); return b.getvalue()


def test_agent_generates_real_image_over_http_then_writes_reference(tmp_path,serve_model):
    calls=[]
    def responder(body,n):
        calls.append(body)
        if body['model']=='image-fixture':
            assert body['n']==1 and body['prompt']=='赛马场景'
            return 200,'application/json',json.dumps({'data':[{'b64_json':base64.b64encode(png()).decode()}],'usage':{'images':1}})
        history=body['messages']
        if history[-1]['role']=='user':
            return sse(tool_reply('image_generate',{'prompt':'赛马场景','path':'assets/scene.png'},'image'))
        if history[-1].get('tool_call_id')=='image':
            result=json.loads(history[-1]['content']);assert result['ok'] and result['data']['generation']=='provider_api'
            return sse(tool_reply('file_write',{'path':'index.html','content':'<img src="assets/scene.png">'},'html'))
        return sse({'role':'assistant','content':'图片及引用文件已经保存。'})
    with serve_model(responder) as (endpoint,_):
        home=configured(tmp_path,endpoint);c=Config(home);c.add_model('images','openai-compatible','image-fixture',endpoint)
        c.data['capabilities']['image']='images';c.save()
        async def scenario():
            rt=Runtime(home)
            try:
                s=rt.store.create();await rt.grant(s['id'],{'mode':'full','write':True,'images':True,'image_limit':1})
                rid=await rt.start(s['id'],'生成场景图片并引用');await rt.wait(rid)
                assert rt.store.run(rid)['status']=='completed'
                root=Path(s['workspace']);assert Image.open(root/'assets/scene.png').size==(48,32)
                assert 'assets/scene.png' in (root/'index.html').read_text()
                assert len([c for c in calls if c['model']=='image-fixture'])==1
                assert any(e['kind']=='media.request' for e in rt.store.events(s['id']))
            finally:await rt.shutdown();rt.close()
        asyncio.run(scenario())


def test_image_budget_stops_new_provider_request(tmp_path):
    count=[]
    class Repeating:
        async def complete(self,*args):return ModelReply(tool_reply('image_generate',{'prompt':'x','path':'x.png'}),{'total_tokens':10},'tool_calls')
    async def scenario():
        rt=Runtime(configured(tmp_path),model_factory=lambda *_:Repeating())
        async def image(args,ctx):count.append(1);return {'path':'x.png'}
        rt.media.image=image
        try:
            s=rt.store.create();await rt.grant(s['id'],{'mode':'full','write':True,'images':True,'image_limit':1})
            rid=await rt.start(s['id'],'生成',max_steps=3);await rt.wait(rid)
            assert len(count)==1
            assert '生图调用上限' in rt.store.history(s['id'])[-1]['content']
        finally:await rt.shutdown();rt.close()
    asyncio.run(scenario())


def test_real_office_files_and_hash_patch(tmp_path):
    from pptx import Presentation
    from docx import Document
    root=tmp_path/'project';root.mkdir();(root/'picture.png').write_bytes(png())
    ctx=ToolContext(root,tmp_path/'state','docs',{'mode':'off'});registry=default_registry()
    async def scenario():
        for extension in ('docx','pptx'):
            result=await registry.execute('document_create',{'path':'lesson.'+extension,'title':'赛马策略','sections':[{'title':'交换次序','body':'先观察，再比较。','bullets':['记录结果','总结策略'],'image':'picture.png'}]},ctx)
            assert result['ok'] and result['data']['editable'] and result['data']['visual_check']=='not_performed'
        slides=Presentation(root/'lesson.pptx').slides
        assert len(slides)==2 and any(s.text=='交换次序' for s in slides[1].shapes if s.has_text_frame)
        assert Document(root/'lesson.docx').paragraphs[0].text=='赛马策略'
        result=await registry.execute('document_read',{'path':'lesson.pptx'},ctx)
        assert '先观察' in result['data']['text']
        await registry.execute('file_write',{'path':'a.txt','content':'hello world'},ctx)
        info=await registry.execute('file_info',{'path':'a.txt'},ctx)
        patch={'path':'a.txt','old':'world','new':'小博','expected_sha256':info['data']['sha256']}
        assert (await registry.execute('file_patch',patch,ctx))['ok']
        assert (root/'a.txt').read_text()=='hello 小博'
        assert not (await registry.execute('file_patch',patch,ctx))['ok']
        manifests=list((tmp_path/'state/versions/docs').glob('*.json'))
        assert len(manifests)==1 and json.loads(manifests[0].read_text())['path']=='a.txt'
    asyncio.run(scenario())


def test_isolated_multifile_preview_and_browser_interaction(tmp_path):
    pytest.importorskip('playwright')
    root=tmp_path/'project';root.mkdir();(root/'assets').mkdir()
    (root/'assets/horse.png').write_bytes(png())
    (root/'index.html').write_text('<!doctype html><title>交互验证</title><link rel="stylesheet" href="style.css"><img src="assets/horse.png"><p id="score">0</p><form id="form"><button id="win" type="submit">获胜</button></form><button id="export">导出</button><script src="game.js"></script>')
    (root/'style.css').write_text('body{background:#f5f4fa;font:20px sans-serif}')
    (root/'game.js').write_text("document.querySelector('#score').textContent=localStorage.getItem('score')||'0';document.querySelector('#form').onsubmit=e=>{e.preventDefault();document.querySelector('#score').textContent='1';localStorage.setItem('score','1')};document.querySelector('#export').onclick=()=>{const a=document.createElement('a');a.href=URL.createObjectURL(new Blob([JSON.stringify({score:1})],{type:'application/json'}));a.download='score.json';a.click()}")
    ctx=ToolContext(root,tmp_path/'state','browser',{'mode':'off'})
    result=asyncio.run(browser_check({'path':'index.html','steps':[{'action':'click','selector':'#win'},{'action':'text_contains','selector':'#score','value':'1'},{'action':'reload','selector':''},{'action':'text_contains','selector':'#score','value':'1'},{'action':'download','selector':'#export','value':'checks/score.json'}],'screenshot':'checks/game.png'},ctx))
    assert result['verified']=='browser_checks_passed',result
    assert result['title']=='交互验证' and len(result['actions'])==5
    assert json.loads((root/'checks/score.json').read_text())=={'score':1}
    assert Image.open(root/'checks/game.png').size==(1280,800)
    allowed=[True];host=PreviewHost()
    try:
        url=host.add(ctx,'index.html',lambda:allowed[0])
        with httpx.Client(trust_env=False) as client:
            assert client.get(url).status_code==200
            assert client.get(url.replace('index.html','assets/horse.png')).content==png()
            assert client.get(host.origin+'/api/bootstrap').status_code==404
            assert client.get(url.replace('index.html','%2e%2e%2fstate/config.json')).status_code==404
            allowed[0]=False;assert client.get(url).status_code==403
    finally:host.close()


def test_remote_asset_url_cannot_target_local_services():
    from lebotclaw_harness.media import download_image
    for url in ('http://example.com/a.png','https://127.0.0.1/a.png','https://user:secret@example.com/a.png'):
        with pytest.raises(ValueError):asyncio.run(download_image(url))
